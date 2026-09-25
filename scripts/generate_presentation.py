"""Generate the official KAIRO Project Purpose presentation matching the app theme."""

from pathlib import Path
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

OUTPUT_FILE = Path("KAIRO_Project_Purpose_Presentation.pptx")

# App Theme Colors
BG_DARK = RGBColor(11, 17, 32)        # #0B1120 deep navy / charcoal
CARD_BG = RGBColor(17, 24, 39)        # #111827 dark slate
CARD_BORDER = RGBColor(30, 41, 59)    # #1E293B slate border
TEAL_PRIMARY = RGBColor(20, 184, 166) # #14B8A6 brand teal
TEAL_LIGHT = RGBColor(45, 212, 191)   # #2DD4BF mint teal
CYAN_ACCENT = RGBColor(56, 189, 248)  # #38BDF8 sky blue
TEXT_WHITE = RGBColor(248, 250, 252)  # #F8FAFC crisp white
TEXT_MUTED = RGBColor(148, 163, 184)  # #94A3B8 slate muted
AMBER_ALERT = RGBColor(245, 158, 11)  # #F59E0B warning amber
GREEN_SUCCESS = RGBColor(16, 185, 129)# #10B981 emerald
PURPLE_ACCENT = RGBColor(99, 102, 241)# #6366F1 indigo


def set_bg(slide, prs):
    """Fill slide with dark app theme background."""
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = BG_DARK
    bg_shape.line.fill.background()
    return bg_shape


def add_header(slide, title: str, category: str, subtitle: str = ""):
    """Add standardized sleek top header with category pill."""
    # Category Pill
    pill = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.5), Inches(2.8), Inches(0.32)
    )
    pill.fill.solid()
    pill.fill.fore_color.rgb = RGBColor(15, 23, 42)
    pill.line.color.rgb = TEAL_PRIMARY
    pill.line.width = Pt(1)
    tf_pill = pill.text_frame
    tf_pill.word_wrap = True
    p_pill = tf_pill.paragraphs[0]
    p_pill.text = category.upper()
    p_pill.font.size = Pt(9)
    p_pill.font.bold = True
    p_pill.font.color.rgb = TEAL_LIGHT
    p_pill.alignment = PP_ALIGN.CENTER

    # Main Title Box
    title_box = slide.shapes.add_textbox(
        Inches(0.8), Inches(0.9), Inches(11.7), Inches(0.7)
    )
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_WHITE

    if subtitle:
        p_sub = tf_title.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.size = Pt(12)
        p_sub.font.color.rgb = TEXT_MUTED


def add_footer(slide, current_idx: int, total_slides: int):
    """Add subtle bottom breadcrumb."""
    footer_box = slide.shapes.add_textbox(
        Inches(0.8), Inches(7.0), Inches(11.7), Inches(0.35)
    )
    tf = footer_box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = f"KAIRO: Temporal Knowledge Graph Based Engineering Intelligence Platform  |  Slide {current_idx} of {total_slides}"
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(71, 85, 105)


def add_card(slide, left, top, width, height, title, items, accent=TEAL_LIGHT, tag=None):
    """Add a structured dark card with header, optional tag, and bullet items."""
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = CARD_BORDER
    card.line.width = Pt(1.2)

    # Accent line top
    accent_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left + Inches(0.2), top + Inches(0.02), width - Inches(0.4), Pt(2.5)
    )
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = accent
    accent_bar.line.fill.background()

    tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), height - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    p_title = tf.paragraphs[0]
    p_title.text = title
    p_title.font.size = Pt(15)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_WHITE

    if tag:
        p_tag = tf.add_paragraph()
        p_tag.text = tag.upper()
        p_tag.font.size = Pt(9)
        p_tag.font.bold = True
        p_tag.font.color.rgb = accent
        p_tag.space_after = Pt(8)

    for item in items:
        p_item = tf.add_paragraph()
        p_item.text = f"- {item}"
        p_item.font.size = Pt(11)
        p_item.font.color.rgb = TEXT_MUTED
        p_item.space_after = Pt(4)


def add_stat_card(slide, left, top, width, height, value, label, subtext, color=TEAL_LIGHT):
    """Add high-impact KPI statistic card."""
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = CARD_BORDER
    card.line.width = Pt(1.2)

    tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    p_val = tf.paragraphs[0]
    p_val.text = value
    p_val.font.size = Pt(32)
    p_val.font.bold = True
    p_val.font.color.rgb = color

    p_lbl = tf.add_paragraph()
    p_lbl.text = label
    p_lbl.font.size = Pt(13)
    p_lbl.font.bold = True
    p_lbl.font.color.rgb = TEXT_WHITE
    p_lbl.space_after = Pt(4)

    p_sub = tf.add_paragraph()
    p_sub.text = subtext
    p_sub.font.size = Pt(10)
    p_sub.font.color.rgb = TEXT_MUTED


def build_presentation():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    total_slides = 10

    # -------------------------------------------------------------
    # SLIDE 1: Title & Purpose Overview
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1, prs)

    # Top Pill
    p1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.0), Inches(4.2), Inches(0.38))
    p1.fill.solid()
    p1.fill.fore_color.rgb = RGBColor(15, 23, 42)
    p1.line.color.rgb = TEAL_PRIMARY
    p1.text_frame.paragraphs[0].text = "ENGINEERING INTELLIGENCE PLATFORM"
    p1.text_frame.paragraphs[0].font.size = Pt(10)
    p1.text_frame.paragraphs[0].font.bold = True
    p1.text_frame.paragraphs[0].font.color.rgb = TEAL_LIGHT
    p1.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Main Title
    tb_title = s1.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(2.2))
    tf_t = tb_title.text_frame
    tf_t.word_wrap = True
    p_h1 = tf_t.paragraphs[0]
    p_h1.text = "KAIRO"
    p_h1.font.size = Pt(56)
    p_h1.font.bold = True
    p_h1.font.color.rgb = TEXT_WHITE

    p_h2 = tf_t.add_paragraph()
    p_h2.text = "A Temporal Knowledge Graph Based Engineering Intelligence Platform\nfor Software Project Reasoning"
    p_h2.font.size = Pt(22)
    p_h2.font.color.rgb = TEAL_LIGHT
    p_h2.space_before = Pt(8)

    p_purpose = tf_t.add_paragraph()
    p_purpose.text = "Project Purpose: Reconstructing and preserving architectural memory across commits, pull requests, and design decisions to eliminate undocumented decision drift."
    p_purpose.font.size = Pt(13)
    p_purpose.font.color.rgb = TEXT_MUTED
    p_purpose.space_before = Pt(14)

    # Bottom Metadata Card
    meta_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.7), Inches(11.7), Inches(1.9))
    meta_card.fill.solid()
    meta_card.fill.fore_color.rgb = CARD_BG
    meta_card.line.color.rgb = CARD_BORDER
    tf_meta = meta_card.text_frame
    tf_meta.word_wrap = True
    tf_meta.margin_left = tf_meta.margin_right = tf_meta.margin_top = tf_meta.margin_bottom = Inches(0.25)

    p_team = tf_meta.paragraphs[0]
    p_team.text = "Team Members: Sharvari Bhondekar, Aman Mehtar, Shruti Gauchandra"
    p_team.font.size = Pt(13)
    p_team.font.bold = True
    p_team.font.color.rgb = TEXT_WHITE

    p_guides = tf_meta.add_paragraph()
    p_guides.text = "Project Mentors: Prof. Kranti Gule, Dr. Tatwadarshi P. Nagarhalli"
    p_guides.font.size = Pt(12)
    p_guides.font.color.rgb = CYAN_ACCENT
    p_guides.space_before = Pt(4)

    p_dept = tf_meta.add_paragraph()
    p_dept.text = "Department of Artificial Intelligence & Data Science, Vidyavardhini's College of Engineering and Technology (VCET)"
    p_dept.font.size = Pt(11)
    p_dept.font.color.rgb = TEXT_MUTED
    p_dept.space_before = Pt(4)

    add_footer(s1, 1, total_slides)

    # -------------------------------------------------------------
    # SLIDE 2: The Core Problem: Engineering Memory Vaporization
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2, prs)
    add_header(s2, "The Critical Problem: Engineering Memory Vaporization", "Industry Challenge", "Why modern software engineering teams lose architectural context over time")

    add_card(
        s2, Inches(0.8), Inches(1.9), Inches(3.64), Inches(4.8),
        "1. Tool Fragmentation",
        [
            "Teams use distributed tools: GitHub, Jira, Slack, wikis, and PR threads.",
            "These systems track the current status of tasks (Done, In Progress), but never preserve why decisions were made.",
            "When original engineers transition or leave, critical institutional memory vanishes permanently.",
            "New engineers inherit codebases with zero verifiable context on legacy trade-offs."
        ],
        accent=PURPLE_ACCENT,
        tag="Siloed Information"
    )

    add_card(
        s2, Inches(4.84), Inches(1.9), Inches(3.64), Inches(4.8),
        "2. Invisible Decision Drift",
        [
            "Documented architecture and live codebases continuously drift apart.",
            "Subsequent pull requests inadvertently violate foundational architectural decisions (ADRs).",
            "Technical debt accumulates silently until massive regressions occur in production.",
            "Teams lack any objective mathematical index to measure or alert on architectural drift."
        ],
        accent=AMBER_ALERT,
        tag="Silent Divergence"
    )

    add_card(
        s2, Inches(8.88), Inches(1.9), Inches(3.64), Inches(4.8),
        "3. Temporal Blindness in AI",
        [
            "Standard Vector RAG retrieves code snippets solely by semantic keyword similarity.",
            "Standard RAG cannot distinguish between an active architectural rule and a superseded choice from 2021.",
            "Results in dangerous hallucinations: AI developer tools recommend obsolete practices as authoritative standards.",
            "Missing bitemporal intervals leads to high risk in AI-assisted coding."
        ],
        accent=CYAN_ACCENT,
        tag="RAG Limitation"
    )

    add_footer(s2, 2, total_slides)

    # -------------------------------------------------------------
    # SLIDE 3: The Project Purpose & Strategic Mission
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3, prs)
    add_header(s3, "Project Purpose: Why KAIRO Exists", "Core Mission", "Building an authoritative, bitemporal engineering knowledge graph for software reasoning")

    # Banner Box
    b_shape = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(11.72), Inches(1.15))
    b_shape.fill.solid()
    b_shape.fill.fore_color.rgb = RGBColor(13, 148, 136)
    b_shape.line.fill.background()
    tf_b = b_shape.text_frame
    tf_b.word_wrap = True
    tf_b.margin_left = tf_b.margin_right = tf_b.margin_top = tf_b.margin_bottom = Inches(0.2)
    p_b1 = tf_b.paragraphs[0]
    p_b1.text = "The Overarching Mission of KAIRO:"
    p_b1.font.size = Pt(11)
    p_b1.font.bold = True
    p_b1.font.color.rgb = RGBColor(204, 251, 241)
    p_b2 = tf_b.add_paragraph()
    p_b2.text = "To transform unstructured developer telemetry into an active, verifiable bitemporal knowledge graph, empowering teams to query, track, and protect the architectural rationale of their codebases."
    p_b2.font.size = Pt(14)
    p_b2.font.bold = True
    p_b2.font.color.rgb = TEXT_WHITE

    # 3 Purpose Pillars
    add_card(
        s3, Inches(0.8), Inches(3.3), Inches(3.64), Inches(3.4),
        "Reconstruct Reasoning",
        [
            "Mine commits, PR discussions, code reviews, and issue threads.",
            "Extract discrete architectural decisions with verbatim provenance quotes.",
            "Enable natural-language queries: 'Why did we choose stateless JWT over Redis sessions?'"
        ],
        accent=TEAL_LIGHT,
        tag="Cognitive Memory"
    )

    add_card(
        s3, Inches(4.84), Inches(3.3), Inches(3.64), Inches(3.4),
        "Pre-Merge Interception",
        [
            "Scan modified files in pending pull requests against the knowledge graph.",
            "Flag conflicts with documented architectural decisions before code merges.",
            "Prevent regressions caused by developer rotation and forgotten standards."
        ],
        accent=AMBER_ALERT,
        tag="Proactive Protection"
    )

    add_card(
        s3, Inches(8.88), Inches(3.3), Inches(3.64), Inches(3.4),
        "Quantify Repository Health",
        [
            "Introduce mathematical metrics for software architecture governance.",
            "Decision Drift Index (DDI) measures codebase divergence.",
            "Engineering Evolution Score (EES) assesses stability, ownership, and maintenance risk."
        ],
        accent=GREEN_SUCCESS,
        tag="Verifiable Governance"
    )

    add_footer(s3, 3, total_slides)

    # -------------------------------------------------------------
    # SLIDE 4: Architectural Innovation: Two-Layer Trust Isolation
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4, prs)
    add_header(s4, "Architectural Innovation: Two-Layer Trust Isolation", "Zero-Hallucination Design", "Separating verifiable deterministic code facts from LLM-interpreted architectural decisions")

    add_card(
        s4, Inches(0.8), Inches(1.9), Inches(5.65), Inches(4.8),
        "Layer 1: Deterministic Fact Graph",
        [
            "Strictly factual engineering telemetry directly synced from GitHub APIs.",
            "Entities: Repositories, Commits, Authors, Pull Requests, Files, Diffs.",
            "Zero LLM involvement: Guarantees 100% mathematical fidelity and zero hallucination risk.",
            "Deterministic Edge Linking: Connects commits to files, authors to PRs, and PRs to issues with precise timestamps.",
            "Foundation of Trust: Serves as the ground truth anchor for all downstream graph walks and impact predictions."
        ],
        accent=CYAN_ACCENT,
        tag="100% Verifiable Reality"
    )

    add_card(
        s4, Inches(6.88), Inches(1.9), Inches(5.65), Inches(4.8),
        "Layer 2: Interpreted Decision Graph",
        [
            "Architectural decisions, trade-offs, and rationale extracted via LLMs.",
            "Bitemporal Validity Intervals: Every decision carries Valid-Time (when it was true in the world) and Transaction-Time (when KAIRO ingested it).",
            "Verbatim Provenance Quotation: Extracted claims must quote exact PR review comments to eliminate fabricated context.",
            "Superseding Graph Edges: When a new decision is ratified (e.g. ADR-048), supersedes edges retire outdated nodes automatically.",
            "Confidence Scoring: Every decision carries an empirical verification score (e.g., 94% confidence)."
        ],
        accent=TEAL_LIGHT,
        tag="Bitemporal LLM Reasoning"
    )

    add_footer(s4, 4, total_slides)

    # -------------------------------------------------------------
    # SLIDE 5: Retrieval Innovation: Three-Mode GraphRAG Engine
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5, prs)
    add_header(s5, "Retrieval Innovation: Three-Mode GraphRAG Engine", "Retrieval Intelligence", "Dynamic mode selection that solves vocabulary mismatch and temporal blindness")

    add_card(
        s5, Inches(0.8), Inches(1.9), Inches(3.64), Inches(4.8),
        "1. Anchored Retrieval",
        [
            "Starts from a specific file path node (e.g. auth/middleware.py).",
            "Performs recursive SQL CTE graph traversal across structural dependencies.",
            "Finds relevant architectural decisions even when they share zero vocabulary with user queries.",
            "Ideal for understanding the full history and intent behind an isolated file or module."
        ],
        accent=TEAL_LIGHT,
        tag="Structural CTE Walk"
    )

    add_card(
        s5, Inches(4.84), Inches(1.9), Inches(3.64), Inches(4.8),
        "2. As-Of Bitemporal Retrieval",
        [
            "Queries the repository state as it existed at any historical snapshot in time.",
            "Example query: 'What was our session management standard in March 2023?'",
            "Filters validity intervals (valid_from <= T <= valid_to) to completely suppress decisions made after timestamp T.",
            "Prevents anachronisms and superseded policies from infecting analysis."
        ],
        accent=PURPLE_ACCENT,
        tag="Time-Bounded Context"
    )

    add_card(
        s5, Inches(8.88), Inches(1.9), Inches(3.64), Inches(4.8),
        "3. Semantic Similarity",
        [
            "High-dimensional vector embedding search powered by pgvector.",
            "Acts as a resilient fallback for broad, conceptual developer queries.",
            "Filters semantic nearest-neighbors through graph connectivity boundaries.",
            "Combines statistical vector search with structural graph precision."
        ],
        accent=CYAN_ACCENT,
        tag="pgvector Fallback"
    )

    add_footer(s5, 5, total_slides)

    # -------------------------------------------------------------
    # SLIDE 6: Novel Quantitative Metrics: DDI & EES
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6, prs)
    add_header(s6, "Novel Quantitative Metrics: DDI & EES", "Architectural Governance", "Mathematical metrics that turn code health and architectural drift into actionable KPIs")

    add_stat_card(
        s6, Inches(0.8), Inches(1.9), Inches(5.65), Inches(2.2),
        "72 / 100",
        "Decision Drift Index (DDI)",
        "Measures the mathematical divergence between documented architectural intent and active code implementation. High DDI indicates elevated architectural risk.",
        color=AMBER_ALERT
    )

    add_stat_card(
        s6, Inches(6.88), Inches(1.9), Inches(5.65), Inches(2.2),
        "85 / 100",
        "Engineering Evolution Score (EES)",
        "Longitudinal repository health index evaluating delivery stability, author ownership diversity, and decision drift over rolling quarterly windows.",
        color=GREEN_SUCCESS
    )

    add_card(
        s6, Inches(0.8), Inches(4.4), Inches(5.65), Inches(2.3),
        "Why DDI Matters for Engineering Teams",
        [
            "Detects silent architecture erosion before it triggers technical debt bankruptcy.",
            "Flags undocumented modifications to foundational layers (auth, database, caching).",
            "Provides automated merge blockers when pull requests spike DDI beyond safe thresholds."
        ],
        accent=AMBER_ALERT
    )

    add_card(
        s6, Inches(6.88), Inches(4.4), Inches(5.65), Inches(2.3),
        "Why EES Matters for Leadership & Evaluators",
        [
            "Replaces subjective code quality reviews with a repeatable empirical score.",
            "Assesses module bus-factor by tracking commit author concentration over time.",
            "Guides sprint refactoring allocation based on factual component decay."
        ],
        accent=TEAL_LIGHT
    )

    add_footer(s6, 6, total_slides)

    # -------------------------------------------------------------
    # SLIDE 7: Pre-Merge Intelligence: Change Impact Scanner
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7, prs)
    add_header(s7, "Pre-Merge Intelligence: Change Impact Scanner", "Pull Request Governance", "Simulated GitHub PR #412 walkthrough: Detecting blast radius and conflicts before merge")

    add_card(
        s7, Inches(0.8), Inches(1.9), Inches(3.64), Inches(4.8),
        "Step 1: Conflict Detection",
        [
            "PR #412 proposes migrating authentication middleware to stateless JWTs (+142, -89).",
            "KAIRO scans modified files and traverses the graph to discover ADR-042.",
            "ADR-042 formally documented that session-based auth was the mandatory standard.",
            "Alert Triggered: 'PR contradicts documented architectural decision ADR-042.'"
        ],
        accent=AMBER_ALERT,
        tag="ADR Collision"
    )

    add_card(
        s7, Inches(4.84), Inches(1.9), Inches(3.64), Inches(4.8),
        "Step 2: Churn & Bug Risk",
        [
            "Scanner queries historical Layer 1 telemetry for auth/middleware.py.",
            "File Churn Telemetry: Modified 47 times in the last 6 months.",
            "Historical Regression Rate: 23% of previous modifications introduced defects.",
            "Predictive Warning: 'Fragile file pattern detected. Additional security review recommended.'"
        ],
        accent=PURPLE_ACCENT,
        tag="Telemetry Risk"
    )

    add_card(
        s7, Inches(8.88), Inches(1.9), Inches(3.64), Inches(4.8),
        "Step 3: Ownership & Clearance",
        [
            "Ownership Analysis: Original architect (@dave) is inactive; @alice holds 67% ownership.",
            "Review Assigned: Automatic ping to active primary maintainers.",
            "Merge Clearance: Team generates superseding ADR-048, updates DDI, and unlocks 96% readiness score.",
            "Architecture remains 100% documented upon merge."
        ],
        accent=GREEN_SUCCESS,
        tag="Resolution"
    )

    add_footer(s7, 7, total_slides)

    # -------------------------------------------------------------
    # SLIDE 8: Agile Work Hierarchy & Planning Poker
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8, prs)
    add_header(s8, "Agile Alignment: Multi-Level Hierarchy & Planning Poker", "Full-Lifecycle Agile", "Connecting high-level business objectives directly to architecture-aware backlog execution")

    add_card(
        s8, Inches(0.8), Inches(1.9), Inches(5.65), Inches(4.8),
        "Multi-Level Work Hierarchy (/demo/hierarchy)",
        [
            "Structured Strategic Alignment: Objectives -> Projects -> Epics -> Issues -> Sub-tasks.",
            "Live Progress Rollups: Automatically calculates point completion percentages at each parent tier.",
            "Architecture Impact Badges: Issues flagged by KAIRO's scanner display conflict warnings directly on the hierarchy tree.",
            "Bidirectional GitHub Synchronization: Syncs issues and milestones with real-time repository commits.",
            "Unified Visibility: Both developers and project managers view the same live architectural roadmap."
        ],
        accent=CYAN_ACCENT,
        tag="Strategic Rollup"
    )

    add_card(
        s8, Inches(6.88), Inches(1.9), Inches(5.65), Inches(4.8),
        "Interactive Planning Poker (/demo/board)",
        [
            "Architecture-Aware Story Point Estimation: Integrated directly into the sprint board.",
            "Fibonacci Estimation Deck: 1, 2, 3, 5, 8, 13, 21 points scale.",
            "Simulated Team Consensus: Real-time vote collation from engineers (@aman, @shruti, @sharvari, @alex).",
            "AI Historical Effort Benchmark: Recommends story point complexity based on historical churn and file fragility.",
            "Instant Backlog Assignment: 1-click commitment updates sprint balances and velocity metrics."
        ],
        accent=TEAL_LIGHT,
        tag="Team Consensus"
    )

    add_footer(s8, 8, total_slides)

    # -------------------------------------------------------------
    # SLIDE 9: Empirical Validation & Technical Results
    # -------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    set_bg(s9, prs)
    add_header(s9, "Empirical Validation & Benchmark Results", "Experimental Findings", "Rigorous benchmarks conducted across open-source repositories and simulated testbeds")

    add_stat_card(
        s9, Inches(0.8), Inches(1.9), Inches(3.64), Inches(2.2),
        "7.1x",
        "Cross-Reference Recovery",
        "Timeline API graph ingestion recovered 7.1x more relationships between PRs, commits, and reviews than standard regex extraction.",
        color=TEAL_LIGHT
    )

    add_stat_card(
        s9, Inches(4.84), Inches(1.9), Inches(3.64), Inches(2.2),
        "56% -> 27%",
        "Overlap Reduction",
        "Dynamic Hub Degree Suppression prevents sweeping commits from dominating search, reducing result-set overlap by more than half.",
        color=CYAN_ACCENT
    )

    add_stat_card(
        s9, Inches(8.88), Inches(1.9), Inches(3.64), Inches(2.2),
        "< 1%",
        "Hallucination Rejection",
        "Verbatim-quote verification on Layer 2 decisions eliminated fabricated citations, achieving a 99%+ provenance precision rate.",
        color=GREEN_SUCCESS
    )

    add_card(
        s9, Inches(0.8), Inches(4.4), Inches(11.72), Inches(2.3),
        "Implementation Architecture & Production Stack",
        [
            "Backend: Python 3.13, FastAPI asynchronous REST API with Pydantic validation.",
            "Database: PostgreSQL with pgvector extension and recursive SQL Common Table Expressions (replacing heavy Neo4j overhead).",
            "Frontend: Next.js 16 with React 19, Tailwind CSS, and Lucide iconography for interactive SaaS exploration.",
            "LLMOps: OpenAI-compatible proxy interface with OpenRouter free-tier LLMs (Cohere, Qwen, Llama) and content-hash caching for reproducible, rate-limit friendly inference."
        ],
        accent=PURPLE_ACCENT,
        tag="Engineering Stack"
    )

    add_footer(s9, 9, total_slides)

    # -------------------------------------------------------------
    # SLIDE 10: Conclusion & Long-Term Impact
    # -------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    set_bg(s10, prs)
    add_header(s10, "Conclusion: The Strategic Impact of KAIRO", "Project Impact", "Preserving intellectual capital and building self-reasoning software engineering ecosystems")

    add_card(
        s10, Inches(0.8), Inches(1.9), Inches(3.64), Inches(4.8),
        "For Engineering Teams",
        [
            "Zero Knowledge Loss: Onboarding engineers immediately understand why code was written the way it is.",
            "Safer Code Reviews: Pull request reviewers know the architectural blast radius before clicking merge.",
            "Confidence in Refactoring: Identify fragile hotspots and outdated decisions before beginning redesigns.",
            "Eliminates guesswork in legacy maintenance."
        ],
        accent=TEAL_LIGHT,
        tag="Developer Productivity"
    )

    add_card(
        s10, Inches(4.84), Inches(1.9), Inches(3.64), Inches(4.8),
        "For Engineering Leadership",
        [
            "Objective Quality Metrics: Track Decision Drift Index (DDI) and Engineering Evolution Score (EES) on executive dashboards.",
            "Eliminate Architectural Debt: Identify divergence months before it causes critical production outages.",
            "Auditability & Compliance: Verbatim provenance provides clear paper-trails for enterprise architecture governance."
        ],
        accent=CYAN_ACCENT,
        tag="Governance & Oversight"
    )

    add_card(
        s10, Inches(8.88), Inches(1.9), Inches(3.64), Inches(4.8),
        "The Research Contribution",
        [
            "Solves Temporal Blindness: Demonstrates bitemporal intervals as the vital solution for LLMOps in codebases.",
            "Two-Layer Trust Paradigm: Proves deterministic telemetry and probabilistic LLM outputs can coexist safely.",
            "A New Category: Moves beyond simple issue trackers toward self-reasoning software engineering memory."
        ],
        accent=GREEN_SUCCESS,
        tag="Academic Innovation"
    )

    add_footer(s10, 10, total_slides)

    # Save presentation
    prs.save(OUTPUT_FILE)
    print(f"Presentation generated successfully: {OUTPUT_FILE}")


if __name__ == "__main__":
    build_presentation()
